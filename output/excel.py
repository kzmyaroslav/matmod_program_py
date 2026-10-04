from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter


THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


def _apply_header_style(row):
    """Оформляет строку заголовков."""

    for cell in row:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="D9EAF7",
        )
        cell.border = THIN_BORDER


def _auto_width(ws):
    """Автоматически подбирает ширину столбцов."""

    for column_cells in ws.columns:
        max_length = 0

        column_letter = get_column_letter(
            column_cells[0].column
        )

        for cell in column_cells:
            if cell.value is not None:
                max_length = max(
                    max_length,
                    len(str(cell.value)),
                )

        ws.column_dimensions[column_letter].width = min(
            max(max_length + 2, 10),
            40,
        )


def _write_matrix(
    ws,
    data,
    row_axis,
    column_axis,
    row_axis_name,
    column_axis_name,
    decimal_places,
):
    """
    Записывает двумерную матрицу освещённости.

    Верхняя строка:
        координаты по X/Y

    Первый столбец:
        координаты по второй оси.
    """

    # Левый верхний угол
    ws.cell(
        row=1,
        column=1,
        value=f"{row_axis_name} \\ {column_axis_name}",
    )

    # Координаты столбцов
    for j, value in enumerate(column_axis, start=2):
        cell = ws.cell(
            row=1,
            column=j,
            value=float(value),
        )
        cell.number_format = "0.000"

    # Координаты строк
    for i, value in enumerate(row_axis, start=2):
        cell = ws.cell(
            row=i,
            column=1,
            value=float(value),
        )
        cell.number_format = "0.000"

    # Значения
    for i in range(data.shape[0]):
        for j in range(data.shape[1]):
            cell = ws.cell(
                row=i + 2,
                column=j + 2,
                value=float(data[i, j]),
            )

            cell.number_format = (
                f"0.{decimal_places * '0'}"
            )

            cell.border = THIN_BORDER
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

    _apply_header_style(ws[1])

    ws.freeze_panes = "B2"

    _auto_width(ws)


def _add_heatmap(
    ws,
    data,
):
    """
    Добавляет цветовую шкалу к таблице.
    """

    start_row = 2
    start_column = 2

    end_row = data.shape[0] + 1
    end_column = data.shape[1] + 1

    start_cell = ws.cell(
        row=start_row,
        column=start_column,
    ).coordinate

    end_cell = ws.cell(
        row=end_row,
        column=end_column,
    ).coordinate

    ws.conditional_formatting.add(
        f"{start_cell}:{end_cell}",
        ColorScaleRule(
            start_type="min",
            start_color="FFF2CC",
            mid_type="percentile",
            mid_value=50,
            mid_color="FFD966",
            end_type="max",
            end_color="C00000",
        ),
    )


def _create_calculation_sheet(
    wb,
    grid,
    illuminance,
    decimal_places,
):
    """Создаёт основной лист расчёта."""

    ws = wb.create_sheet("Расчет")

    axis_names = {
        "X": ("Y, м", "Z, м"),
        "Y": ("X, м", "Z, м"),
        "Z": ("X, м", "Y, м"),
    }

    row_axis_name, column_axis_name = axis_names[
        grid.plane
    ]

    _write_matrix(
        ws=ws,
        data=illuminance,
        row_axis=grid.row_axis,
        column_axis=grid.column_axis,
        row_axis_name=row_axis_name,
        column_axis_name=column_axis_name,
        decimal_places=decimal_places,
    )

    _add_heatmap(
        ws,
        illuminance,
    )

    ws.sheet_view.showGridLines = False

    return ws


def _create_source_sheet(
    wb,
    source,
    contribution,
    grid,
    decimal_places,
):
    """Создаёт лист с вкладом конкретного источника."""

    ws = wb.create_sheet(
        f"Источник_{source.index}"
    )

    axis_names = {
        "X": ("Y, м", "Z, м"),
        "Y": ("X, м", "Z, м"),
        "Z": ("X, м", "Y, м"),
    }

    row_axis_name, column_axis_name = axis_names[
        grid.plane
    ]

    _write_matrix(
        ws=ws,
        data=contribution,
        row_axis=grid.row_axis,
        column_axis=grid.column_axis,
        row_axis_name=row_axis_name,
        column_axis_name=column_axis_name,
        decimal_places=decimal_places,
    )

    _add_heatmap(
        ws,
        contribution,
    )

    ws.sheet_view.showGridLines = False

    return ws


def _create_statistics_sheet(
    wb,
    statistics,
    decimal_places,
):
    """Создаёт лист статистики."""

    ws = wb.create_sheet("Статистика")

    ws["A1"] = "Показатель"
    ws["B1"] = "Значение"

    _apply_header_style(ws[1])

    rows = [
        ("Минимальная освещённость Emin, лк", statistics["Emin"]),
        ("Максимальная освещённость Emax, лк", statistics["Emax"]),
        ("Средняя освещённость Eavg, лк", statistics["Eavg"]),
        ("Emin / Emax", statistics["Emin/Emax"]),
        ("Eavg / Emax", statistics["Eavg/Emax"]),
    ]

    for row_number, (name, value) in enumerate(
        rows,
        start=2,
    ):
        ws.cell(
            row=row_number,
            column=1,
            value=name,
        )

        ws.cell(
            row=row_number,
            column=2,
            value=float(value),
        )

        ws.cell(
            row=row_number,
            column=1,
        ).border = THIN_BORDER

        ws.cell(
            row=row_number,
            column=2,
        ).border = THIN_BORDER

    for row in range(2, 5):
        ws.cell(
            row=row,
            column=2,
        ).number_format = (
            f"0.{decimal_places * '0'}"
        )

    for row in range(5, 7):
        ws.cell(
            row=row,
            column=2,
        ).number_format = "0.000"

    ws.sheet_view.showGridLines = False

    _auto_width(ws)

    return ws


def _create_parameters_sheet(
    wb,
    room,
    grid_config,
    grid,
    calculation_config,
):
    """Создаёт лист с параметрами расчёта."""

    ws = wb.create_sheet("Параметры")

    ws["A1"] = "Параметр"
    ws["B1"] = "Значение"

    _apply_header_style(ws[1])

    parameters = [
        ("Шаг сетки, м", grid_config.step),
        ("Расчётная плоскость", grid.plane),
        ("Координата плоскости, м", grid.coordinate),
        ("Отладочный режим", calculation_config.debug),
        (
            "Знаков после запятой",
            calculation_config.decimal_places,
        ),
    ]

    for row_number, (name, value) in enumerate(
        parameters,
        start=2,
    ):
        ws.cell(
            row=row_number,
            column=1,
            value=name,
        )

        ws.cell(
            row=row_number,
            column=2,
            value=value,
        )

    ws.sheet_view.showGridLines = False

    _auto_width(ws)

    return ws


def _create_room_sheet(
    wb,
    room,
):
    """Создаёт лист с параметрами помещения."""

    ws = wb.create_sheet("Помещение")

    ws["A1"] = "Размер"
    ws["B1"] = "Значение, м"

    _apply_header_style(ws[1])

    dimensions = [
        ("X", room.x),
        ("Y", room.y),
        ("Z", room.z),
    ]

    for row_number, (name, value) in enumerate(
        dimensions,
        start=2,
    ):
        ws.cell(
            row=row_number,
            column=1,
            value=name,
        )

        ws.cell(
            row=row_number,
            column=2,
            value=float(value),
        )

        ws.cell(
            row=row_number,
            column=2,
        ).number_format = "0.000"

    ws.sheet_view.showGridLines = False

    _auto_width(ws)

    return ws


def _create_source_parameters_sheet(
    wb,
    source,
):
    """
    Создаёт отдельный лист с параметрами источника.
    """

    ws = wb.create_sheet(
        f"Источник_{source.index}_параметры"
    )

    ws["A1"] = "Параметр"
    ws["B1"] = "Значение"

    _apply_header_style(ws[1])

    parameters = source.parameter_dict()

    for row_number, (name, value) in enumerate(
        parameters.items(),
        start=2,
    ):
        ws.cell(
            row=row_number,
            column=1,
            value=name,
        )

        ws.cell(
            row=row_number,
            column=2,
            value=value,
        )

    ws.sheet_view.showGridLines = False

    _auto_width(ws)

    return ws


def export_to_excel(
    output_path,
    room,
    grid_config,
    grid,
    sources,
    illuminance,
    contributions,
    statistics,
    calculation_config,
):
    """
    Полностью формирует Excel-файл с результатами расчёта.

    Parameters
    ----------
    output_path : str | Path
        Путь к Excel-файлу.

    room : Room
        Помещение.

    grid_config : GridConfig
        Настройки сетки.

    grid : GridData
        Расчётная сетка.

    sources : list[LightSource]
        Источники.

    illuminance : np.ndarray
        Суммарная освещённость.

    contributions : dict
        Вклад каждого источника.

    statistics : dict
        Статистика.

    calculation_config : CalculationConfig
        Настройки расчёта.
    """

    output_path = Path(output_path)

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    wb = Workbook()

    # Удаляем автоматически созданный лист.
    default_sheet = wb.active
    wb.remove(default_sheet)

    decimal_places = (
        calculation_config.decimal_places
    )

    # Основной расчёт
    _create_calculation_sheet(
        wb,
        grid,
        illuminance,
        decimal_places,
    )

    # Статистика
    _create_statistics_sheet(
        wb,
        statistics,
        decimal_places,
    )

    # Вклады источников
    for source in sources:

        contribution = contributions[
            source.index
        ]

        _create_source_sheet(
            wb=wb,
            source=source,
            contribution=contribution,
            grid=grid,
            decimal_places=decimal_places,
        )

        _create_source_parameters_sheet(
            wb,
            source,
        )

    # Параметры
    _create_parameters_sheet(
        wb,
        room,
        grid_config,
        grid,
        calculation_config,
    )

    # Помещение
    _create_room_sheet(
        wb,
        room,
    )

    # Активируем основной лист
    wb.active = wb.sheetnames.index(
        "Расчет"
    )

    wb.save(output_path)

    return output_path